#!/usr/bin/env python3
"""Arma el informe final en .docx a partir de los capítulos en Markdown y la plantilla del avance.

Uso (desde la raíz del proyecto, con el venv activado y python-docx instalado):
    pip install -r requirements-informe.txt
    python docs/informe/armar_docx.py

Lee docs/informe/0*.md en orden, conserva la portada de docs/informe_avance_original.docx y vuelca
el resto con los estilos de esa plantilla (Arial, títulos, epígrafes, Roboto Mono para el código).
Escribe docs/informe/TFI_Noguerol_2026.docx. El índice queda como campo de Word: hay que
actualizarlo al abrir el documento (en Google Docs se regenera desde el propio índice).
"""
import copy
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Emu, RGBColor
from PIL import Image

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(os.path.dirname(AQUI))
PLANTILLA = os.path.join(RAIZ, "docs", "informe_avance_original.docx")
SALIDA = os.path.join(AQUI, "TFI_Noguerol_2026.docx")
FIGURA_MERMAID = os.path.join(AQUI, "img", "figura_4_1_arquitectura.png")

# Medidas de la plantilla: A4 con márgenes de 1440 twips -> 9029 twips de ancho útil
ANCHO_UTIL_TWIPS = 11909 - 2 * 1440
ANCHO_UTIL_EMU = Emu(ANCHO_UTIL_TWIPS * 635)
ALTO_MAX_FIGURA_EMU = Emu(int(8.3 * 914400))

FUENTE_CODIGO = "Roboto Mono"
GRIS_ENCABEZADO = "D9D9D9"
GRIS_CODIGO = "F2F2F2"
GRIS_BORDE = "808080"

# Cantidad de párrafos de la portada en la plantilla (logo, carrera, título, autor, lugar y año)
PARRAFOS_PORTADA = 14

# Índices de numbering.xml de la plantilla: 1 es viñetas, 6 es decimal
ABSTRACT_VINETAS = 1
ABSTRACT_DECIMAL = 6


# ---------------------------------------------------------------- utilidades de XML

# Orden que exige el esquema OOXML para los hijos de pPr y de tblPr (solo los que se usan acá)
ORDEN_PPR = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "numPr", "pBdr", "shd", "tabs",
             "spacing", "ind", "jc", "rPr"]
ORDEN_TBLPR = ["tblStyle", "tblW", "jc", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]


def _insertar_ordenado(padre, hijo, orden):
    """Inserta hijo en padre respetando el orden del esquema: antes del primer elemento que lo sigue."""
    nombre = hijo.tag.split("}")[1]
    sucesores = {qn(f"w:{n}") for n in orden[orden.index(nombre) + 1:]}
    for existente in padre:
        if existente.tag in sucesores:
            existente.addprevious(hijo)
            return
    padre.append(hijo)


def _set_spacing(p, antes=240, despues=240, linea=None):
    pf = p.paragraph_format
    pf.space_before = Pt(antes / 20)
    pf.space_after = Pt(despues / 20)
    if linea is not None:
        pf.line_spacing = linea


def _sombrear(elemento, color):
    """Fondo de color para un párrafo (pPr) o una celda (tcPr)."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    if elemento.tag == qn("w:pPr"):
        _insertar_ordenado(elemento, shd, ORDEN_PPR)
    else:
        elemento.append(shd)


def _fuente(run, nombre=None, tamano=None, negrita=None, cursiva=None, color=None):
    if nombre:
        run.font.name = nombre
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for atributo in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(atributo), nombre)
    if tamano:
        run.font.size = Pt(tamano)
    if negrita is not None:
        run.font.bold = negrita
    if cursiva is not None:
        run.font.italic = cursiva
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _salto_de_pagina_antes(p):
    p.paragraph_format.page_break_before = True


def _mantener_con_siguiente(p):
    p.paragraph_format.keep_with_next = True


class Numeracion:
    """Crea instancias de lista nuevas (w:num) para que cada lista numerada arranque en 1."""

    def __init__(self, doc):
        self.numbering = doc.part.numbering_part.element
        ids = [int(n.get(qn("w:numId"))) for n in self.numbering.findall(qn("w:num"))]
        self.siguiente = max(ids) + 1

    def nueva(self, abstract_id):
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(self.siguiente))
        abstracto = OxmlElement("w:abstractNumId")
        abstracto.set(qn("w:val"), str(abstract_id))
        num.append(abstracto)
        # Reinicia la cuenta en 1 aunque el abstractNum se haya usado antes
        override = OxmlElement("w:lvlOverride")
        override.set(qn("w:ilvl"), "0")
        inicio = OxmlElement("w:startOverride")
        inicio.set(qn("w:val"), "1")
        override.append(inicio)
        num.append(override)
        self.numbering.append(num)
        self.siguiente += 1
        return self.siguiente - 1


def _aplicar_lista(p, num_id, nivel=0):
    ppr = p._element.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(nivel))
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(num_id))
    numpr.append(ilvl)
    numpr.append(numid)
    _insertar_ordenado(ppr, numpr, ORDEN_PPR)
    p.paragraph_format.left_indent = Pt(36 * (nivel + 1))
    p.paragraph_format.first_line_indent = Pt(-18)


def _agregar_hipervinculo(p, url, texto=None, tamano=12):
    """Agrega un hipervínculo al párrafo (python-docx no lo trae de fábrica)."""
    r_id = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                            is_external=True)
    enlace = OxmlElement("w:hyperlink")
    enlace.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1155CC")
    sub = OxmlElement("w:u")
    sub.set(qn("w:val"), "single")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(tamano * 2)))
    rpr.append(color)
    rpr.append(sub)
    rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = texto or url
    run.append(t)
    enlace.append(run)
    p._element.append(enlace)


# ---------------------------------------------------------------- texto con formato en línea

# Orden: código, enlace con texto, URL entre <>, negrita, cursiva
PATRON_INLINE = re.compile(
    r"(`[^`]+`)|(\[[^\]]+\]\(https?://[^)]+\))|(<https?://[^>]+>)|(\*\*[^*]+\*\*)|(\*[^*]+\*)"
)


def limpiar_markdown(texto):
    """Saca las marcas de formato (para títulos e índice)."""
    return re.sub(r"[`*]", "", texto)


def escribir_inline(p, texto, tamano=12, negrita_base=False, cursiva_base=False, tamano_codigo=None):
    """Escribe texto con `código`, **negrita**, *cursiva* y enlaces en el párrafo."""
    tamano_codigo = tamano_codigo or max(tamano - 1, 7)
    pos = 0
    for m in PATRON_INLINE.finditer(texto):
        if m.start() > pos:
            run = p.add_run(texto[pos:m.start()])
            _fuente(run, tamano=tamano, negrita=negrita_base or None, cursiva=cursiva_base or None)
        trozo = m.group(0)
        if m.group(1):
            run = p.add_run(trozo[1:-1])
            _fuente(run, nombre=FUENTE_CODIGO, tamano=tamano_codigo, negrita=negrita_base or None)
        elif m.group(2):
            visible, url = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", trozo).groups()
            _agregar_hipervinculo(p, url, visible, tamano)
        elif m.group(3):
            _agregar_hipervinculo(p, trozo[1:-1], tamano=tamano)
        elif m.group(4):
            run = p.add_run(trozo[2:-2])
            _fuente(run, tamano=tamano, negrita=True, cursiva=cursiva_base or None)
        else:
            run = p.add_run(trozo[1:-1])
            _fuente(run, tamano=tamano, cursiva=True, negrita=negrita_base or None)
        pos = m.end()
    if pos < len(texto):
        run = p.add_run(texto[pos:])
        _fuente(run, tamano=tamano, negrita=negrita_base or None, cursiva=cursiva_base or None)


# ---------------------------------------------------------------- bloques de Markdown

def leer_bloques(ruta):
    """Parser mínimo de Markdown: devuelve una lista de (tipo, contenido) para lo que usan los capítulos."""
    lineas = open(ruta, encoding="utf-8").read().split("\n")
    bloques = []
    i = 0
    while i < len(lineas):
        linea = lineas[i]
        if not linea.strip():
            i += 1
            continue
        if linea.startswith("```"):
            lenguaje = linea[3:].strip()
            cuerpo = []
            i += 1
            while i < len(lineas) and not lineas[i].startswith("```"):
                cuerpo.append(lineas[i])
                i += 1
            i += 1
            bloques.append(("codigo", (lenguaje, cuerpo)))
            continue
        m = re.match(r"^(#{1,3}) (.*)$", linea)
        if m:
            bloques.append(("titulo", (len(m.group(1)), m.group(2).strip())))
            i += 1
            continue
        if linea.startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].startswith("|"):
                celdas = [c.strip() for c in lineas[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in celdas):
                    filas.append(celdas)
                i += 1
            bloques.append(("tabla", filas))
            continue
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", linea)
        if m:
            bloques.append(("imagen", os.path.normpath(os.path.join(os.path.dirname(ruta), m.group(2)))))
            i += 1
            continue
        m = re.match(r"^(-|\d+\.) (.*)$", linea)
        if m:
            numerada = m.group(1) != "-"
            items = []
            while i < len(lineas):
                m = re.match(r"^(-|\d+\.) (.*)$", lineas[i])
                if m and (m.group(1) != "-") == numerada:
                    items.append(m.group(2))
                    i += 1
                    # Continuaciones del ítem: líneas sangradas
                    while i < len(lineas) and lineas[i].startswith("  ") and lineas[i].strip():
                        items[-1] += " " + lineas[i].strip()
                        i += 1
                else:
                    break
            bloques.append(("lista", (numerada, items)))
            continue
        # Párrafo: líneas seguidas hasta una en blanco o el inicio de otro bloque
        parrafo = []
        while i < len(lineas) and lineas[i].strip() and not re.match(r"^(#{1,3} |```|\||!\[|- |\d+\. )", lineas[i]):
            parrafo.append(lineas[i].strip())
            i += 1
        texto = " ".join(parrafo)
        if re.match(r"^\*\*(Figura|Tabla) [\dA-Z]+\.\d+\.\*\*", texto):
            bloques.append(("epigrafe", texto))
        else:
            bloques.append(("parrafo", texto))
    return bloques


# ---------------------------------------------------------------- escritura en el documento

class Armador:
    def __init__(self, doc):
        self.doc = doc
        self.numeracion = Numeracion(doc)
        self.titulos = []  # (nivel, texto) para el índice
        self.advertencias = []

    # --- párrafos

    def parrafo(self, texto, tamano=12, alineacion=WD_ALIGN_PARAGRAPH.JUSTIFY, antes=240, despues=240):
        p = self.doc.add_paragraph()
        p.alignment = alineacion
        _set_spacing(p, antes, despues)
        escribir_inline(p, texto, tamano)
        return p

    def titulo(self, nivel, texto):
        texto_limpio = limpiar_markdown(texto)
        self.titulos.append((nivel, texto_limpio))
        p = self.doc.add_paragraph(style=f"Heading {nivel}")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if nivel > 1 else WD_ALIGN_PARAGRAPH.LEFT
        if nivel == 1:
            _salto_de_pagina_antes(p)
            _set_spacing(p, 280, 240)
        else:
            _set_spacing(p, 240, 240)
        escribir_inline(p, texto, tamano=None, negrita_base=True,
                        tamano_codigo={1: 18, 2: 15, 3: 13}[nivel])
        return p

    def epigrafe(self, texto, antes=120, despues=240):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_spacing(p, antes, despues)
        escribir_inline(p, texto, tamano=10, cursiva_base=True)
        return p

    def lista(self, numerada, items):
        num_id = self.numeracion.nueva(ABSTRACT_DECIMAL if numerada else ABSTRACT_VINETAS)
        for k, item in enumerate(items):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _set_spacing(p, 0, 120 if k < len(items) - 1 else 240)
            _aplicar_lista(p, num_id)
            escribir_inline(p, item, 12)

    def referencia(self, texto):
        """Entrada de la lista de referencias: sangría francesa, alineada a la izquierda."""
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _set_spacing(p, 0, 240)
        p.paragraph_format.left_indent = Pt(36)
        p.paragraph_format.first_line_indent = Pt(-36)
        # Las URL sueltas se vuelven enlaces
        partes = re.split(r"(https?://\S+)", texto)
        for parte in partes:
            if parte.startswith("http"):
                _agregar_hipervinculo(p, parte)
            elif parte:
                escribir_inline(p, parte, 12)

    # --- código

    def codigo(self, lenguaje, lineas):
        if lenguaje == "mermaid":
            self.imagen(FIGURA_MERMAID)
            return
        while lineas and not lineas[-1].strip():
            lineas.pop()
        tamano = 9 if max((len(l) for l in lineas), default=0) <= 100 else 8
        for k, linea in enumerate(lineas):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            _set_spacing(p, 0, 0, 1.0)
            p.paragraph_format.left_indent = Pt(8)
            p.paragraph_format.right_indent = Pt(8)
            if k < len(lineas) - 1:
                _mantener_con_siguiente(p)
            _sombrear(p._element.get_or_add_pPr(), GRIS_CODIGO)
            run = p.add_run(linea.replace("\t", "    ") or " ")
            _fuente(run, nombre=FUENTE_CODIGO, tamano=tamano)
        # Separación después del bloque
        p.paragraph_format.space_after = Pt(12)

    # --- figuras

    def imagen(self, ruta):
        if not os.path.isfile(ruta):
            self.advertencias.append(f"Falta la imagen {ruta}")
            p = self.parrafo(f"[FALTA LA IMAGEN: {os.path.basename(ruta)}]", alineacion=WD_ALIGN_PARAGRAPH.CENTER)
            _mantener_con_siguiente(p)
            return
        with Image.open(ruta) as im:
            ancho_px, alto_px = im.size
        ancho = ANCHO_UTIL_EMU
        alto = int(ancho * alto_px / ancho_px)
        if alto > ALTO_MAX_FIGURA_EMU:
            alto = ALTO_MAX_FIGURA_EMU
            ancho = int(alto * ancho_px / alto_px)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_spacing(p, 240, 0)
        _mantener_con_siguiente(p)
        p.add_run().add_picture(ruta, width=Emu(ancho), height=Emu(alto))

    # --- tablas

    def tabla(self, filas):
        columnas = max(len(f) for f in filas)
        filas = [f + [""] * (columnas - len(f)) for f in filas]
        largo_maximo = max(sum(len(c) for c in f) for f in filas)
        if columnas >= 6 or largo_maximo > 400:
            tamano = 8
        elif columnas >= 4 or largo_maximo > 250:
            tamano = 9
        else:
            tamano = 10

        tabla = self.doc.add_table(rows=len(filas), cols=columnas)
        tabla.alignment = 1  # centrada
        tabla.autofit = False
        tblpr = tabla._tbl.tblPr

        # Bordes finos en toda la tabla
        bordes = OxmlElement("w:tblBorders")
        for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{lado}")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "4")
            b.set(qn("w:space"), "0")
            b.set(qn("w:color"), GRIS_BORDE)
            bordes.append(b)
        _insertar_ordenado(tblpr, bordes, ORDEN_TBLPR)

        # Ancho total = ancho útil; cada columna proporcional al largo de su contenido, con un piso
        anchos_contenido = [max(12, min(60, max(len(f[c]) for f in filas))) for c in range(columnas)]
        total = sum(anchos_contenido)
        anchos = [int(ANCHO_UTIL_TWIPS * a / total) for a in anchos_contenido]
        tblw = tblpr.find(qn("w:tblW"))
        if tblw is None:
            tblw = OxmlElement("w:tblW")
            tblpr.append(tblw)
        tblw.set(qn("w:type"), "dxa")
        tblw.set(qn("w:w"), str(sum(anchos)))
        for c, gridcol in enumerate(tabla._tbl.tblGrid.findall(qn("w:gridCol"))):
            gridcol.set(qn("w:w"), str(anchos[c]))

        for r, fila in enumerate(filas):
            fila_xml = tabla.rows[r]._tr
            trpr = fila_xml.get_or_add_trPr()
            cant_split = OxmlElement("w:cantSplit")
            trpr.append(cant_split)
            if r == 0:
                encabezado = OxmlElement("w:tblHeader")
                trpr.append(encabezado)
            for c, contenido in enumerate(fila):
                celda = tabla.cell(r, c)
                celda.width = Emu(anchos[c] * 635)
                tcpr = celda._tc.get_or_add_tcPr()
                if r == 0:
                    _sombrear(tcpr, GRIS_ENCABEZADO)
                p = celda.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                _set_spacing(p, 20, 20, 1.0)
                escribir_inline(p, contenido, tamano, negrita_base=(r == 0), tamano_codigo=tamano - 1)
        # Aire después de la tabla
        p = self.doc.add_paragraph()
        _set_spacing(p, 0, 120)

    # --- índice

    def indice(self):
        """Campo TOC de Word con los títulos ya volcados como resultado provisorio (sin número de página)."""
        p = self.doc.add_paragraph()
        _salto_de_pagina_antes(p)
        _set_spacing(p, 0, 240)
        run = p.add_run("Índice")
        _fuente(run, tamano=14, negrita=True)

        def fld(tipo):
            r = OxmlElement("w:r")
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), tipo)
            r.append(f)
            return r

        primero = self.doc.add_paragraph()
        _set_spacing(primero, 0, 0)
        primero._element.append(fld("begin"))
        r = OxmlElement("w:r")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = ' TOC \\o "1-3" \\h \\z \\u '
        r.append(instr)
        primero._element.append(r)
        primero._element.append(fld("separate"))
        run = primero.add_run("(Índice provisorio: actualizar el campo para que aparezcan los números de página.)")
        _fuente(run, tamano=9, cursiva=True)

        ultimo = primero
        for nivel, texto in self.titulos:
            p = self.doc.add_paragraph()
            _set_spacing(p, 0, 40)
            p.paragraph_format.left_indent = Pt(18 * (nivel - 1))
            run = p.add_run(texto)
            _fuente(run, tamano=11, negrita=(nivel == 1))
            ultimo = p
        ultimo._element.append(fld("end"))


# ---------------------------------------------------------------- armado del documento

def conservar_portada(doc):
    """Deja en el cuerpo solo los párrafos de la portada de la plantilla (y las propiedades de sección)."""
    body = doc.element.body
    hijos = list(body)
    parrafos_vistos = 0
    for hijo in hijos:
        if hijo.tag == qn("w:sectPr"):
            continue
        if hijo.tag == qn("w:p") and parrafos_vistos < PARRAFOS_PORTADA:
            parrafos_vistos += 1
            continue
        body.remove(hijo)


def volcar_resumen(armador, ruta):
    """El Resumen va como en el avance: título centrado (sin estilo de capítulo) y texto con sangría."""
    doc = armador.doc
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _salto_de_pagina_antes(p)
    _set_spacing(p, 0, 240)
    run = p.add_run("Resumen")
    _fuente(run, tamano=14, negrita=True)
    for tipo, contenido in leer_bloques(ruta):
        if tipo == "titulo":
            continue
        if tipo != "parrafo":
            raise ValueError(f"El Resumen solo admite párrafos; apareció {tipo}")
        p = armador.parrafo(contenido)
        p.paragraph_format.left_indent = Pt(30)
        p.paragraph_format.right_indent = Pt(30)


def volcar_capitulo(armador, ruta):
    es_referencias = os.path.basename(ruta).startswith("08_")
    for tipo, contenido in leer_bloques(ruta):
        if tipo == "titulo":
            armador.titulo(*contenido)
        elif tipo == "parrafo":
            if es_referencias and re.search(r"\(\d{4}\)|\(s\.f\.", contenido):
                armador.referencia(contenido)
            else:
                armador.parrafo(contenido)
        elif tipo == "epigrafe":
            armador.epigrafe(contenido)
        elif tipo == "lista":
            armador.lista(*contenido)
        elif tipo == "codigo":
            armador.codigo(*contenido)
        elif tipo == "imagen":
            armador.imagen(contenido)
        elif tipo == "tabla":
            armador.tabla(contenido)
        else:
            raise ValueError(tipo)


def main():
    capitulos = sorted(f for f in os.listdir(AQUI) if re.match(r"^\d\d_.*\.md$", f))
    resumen = [c for c in capitulos if c.startswith("00_")]
    cuerpo = [c for c in capitulos if not c.startswith("00_")]

    doc = Document(PLANTILLA)
    conservar_portada(doc)
    armador = Armador(doc)

    for c in resumen:
        volcar_resumen(armador, os.path.join(AQUI, c))

    # El índice tiene que ir antes de los capítulos, pero necesita sus títulos: se arma al final
    # en un documento aparte y se mueve. Más simple: se recorren los capítulos dos veces.
    for c in cuerpo:
        for tipo, contenido in leer_bloques(os.path.join(AQUI, c)):
            if tipo == "titulo":
                armador.titulos.append((contenido[0], limpiar_markdown(contenido[1])))
    armador.indice()
    armador.titulos = []

    for c in cuerpo:
        volcar_capitulo(armador, os.path.join(AQUI, c))

    doc.save(SALIDA)
    print(f"Escrito {os.path.relpath(SALIDA, RAIZ)}")
    print(f"  capítulos: {len(cuerpo)} + resumen, títulos en el índice: {len(armador.titulos)}, "
          f"tablas: {len(doc.tables)}, imágenes: {len(doc.inline_shapes)}")
    for a in armador.advertencias:
        print(f"  [!] {a}")
    return 1 if armador.advertencias else 0


if __name__ == "__main__":
    sys.exit(main())
